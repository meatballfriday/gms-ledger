"""Generate the course's spreadsheet templates (BraceWorks example figures, fictional)."""
import os
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter as L
from openpyxl.formatting.rule import FormulaRule, CellIsRule
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.comments import Comment

OUT = '/home/claude/course/templates'
os.makedirs(OUT, exist_ok=True)

F = 'Arial'
f_norm = Font(name=F, size=10)
f_bold = Font(name=F, size=10, bold=True)
f_title = Font(name=F, size=14, bold=True)
f_sub = Font(name=F, size=10, italic=True, color='555555')
f_in = Font(name=F, size=10, color='0000FF')
f_link = Font(name=F, size=10, color='008000')
f_head = Font(name=F, size=10, bold=True, color='FFFFFF')
fill_in = PatternFill('solid', fgColor='FFFF00')
fill_head = PatternFill('solid', fgColor='0C6A61')
fill_band = PatternFill('solid', fgColor='EEF3F1')
fill_red = PatternFill('solid', fgColor='F8D0CB')
fill_amber = PatternFill('solid', fgColor='FCE8C3')
fill_green = PatternFill('solid', fgColor='D5ECD9')
thin = Side(style='thin', color='B0BEB8')
top_line = Border(top=Side(style='thin', color='16232A'))
total_line = Border(top=Side(style='thin', color='16232A'), bottom=Side(style='double', color='16232A'))
USD = '$#,##0;($#,##0);"-"'
USD2 = '$#,##0.00;($#,##0.00);"-"'
PCT = '0.0%;(0.0%);"-"'
NUM = '#,##0;(#,##0);"-"'
MULT = '0.00"x"'


def base(ws):
    for row in ws.iter_rows():
        for c in row:
            if c.font == Font():
                c.font = f_norm


def setw(ws, widths):
    for k, v in widths.items():
        ws.column_dimensions[k].width = v


def inp(c, v, fmt=None):
    c.value = v
    c.font = f_in
    c.fill = fill_in
    if fmt:
        c.number_format = fmt


def head_row(ws, r, labels, start=1):
    for i, t in enumerate(labels):
        c = ws.cell(r, start + i, t)
        c.font = f_head
        c.fill = fill_head
        c.alignment = Alignment(horizontal='center' if i else 'left', vertical='center', wrap_text=True)


def instructions(wb, title, purpose, steps, extra=None):
    ws = wb.create_sheet('How to use', 0)
    ws['A1'] = title; ws['A1'].font = f_title
    ws['A2'] = 'The GM\'s Ledger course template. Example figures are for BraceWorks Medical, a fictional company: replace them with your own.'; ws['A2'].font = f_sub
    ws['A4'] = 'What it\'s for'; ws['A4'].font = f_bold
    ws['A5'] = purpose; ws['A5'].alignment = Alignment(wrap_text=True, vertical='top'); ws.row_dimensions[5].height = 45
    ws['A7'] = 'How to use it'; ws['A7'].font = f_bold
    r = 8
    for i, s in enumerate(steps, 1):
        ws.cell(r, 1, f'{i}. {s}').alignment = Alignment(wrap_text=True, vertical='top'); ws.row_dimensions[r].height = 30; r += 1
    r += 1
    ws.cell(r, 1, 'Colour legend').font = f_bold; r += 1
    c = ws.cell(r, 1, 'Blue text on yellow: an input. Type your own figures here.'); c.font = f_in; c.fill = fill_in; r += 1
    ws.cell(r, 1, 'Black text: a formula. Don\'t overwrite these.'); r += 1
    c = ws.cell(r, 1, 'Green text: a figure brought in from another sheet.'); c.font = f_link; r += 1
    if extra:
        r += 1
        for e in extra:
            ws.cell(r, 1, e).alignment = Alignment(wrap_text=True, vertical='top'); ws.row_dimensions[r].height = 30; r += 1
    ws.column_dimensions['A'].width = 110
    base(ws)
    return ws


# ------------------------------------------------------------------ 1. 13-week cash forecast
def cash13():
    wb = Workbook(); ws = wb.active; ws.title = 'Forecast'
    ws['A1'] = '13-week cash forecast'; ws['A1'].font = f_title
    ws['A2'] = 'All figures in US dollars. Update every week and roll forward one week.'; ws['A2'].font = f_sub
    ws['A3'] = 'First week ending (date)'; inp(ws['B3'], '2026-04-03'); ws['B3'].number_format = '@'
    ws['B3'].comment = Comment('Type the date the first week ends, e.g. 2026-04-03.', 'Template')
    head_row(ws, 4, ['US dollars'] + [f'Week {i}' for i in range(1, 14)] + ['Total'])
    ws['A5'] = 'Opening net cash (cash minus overdraft drawn)'; ws['A5'].font = f_bold
    inp(ws['B5'], -59400, USD)
    ws['B5'].comment = Comment('Week 1 only: cash in the bank minus overdraft used. Example: $40,600 cash − $100,000 overdraft = ($59,400).', 'Template')
    for col in range(3, 15):
        ws.cell(5, col, f'={L(col-1)}26').number_format = USD
    ws['A6'] = 'Receipts'; ws['A6'].font = f_bold
    rec = [('US distributor A', [0, 0, 25000, 0, 0, 0, 25000, 0, 0, 0, 25000, 0, 0]),
           ('Australian hospital supplier', [6000, 0, 6000, 6000, 0, 6000, 6000, 0, 6000, 6000, 0, 6000, 6000]),
           ('German distributor', [8000, 0, 5000, 0, 5000, 0, 5000, 0, 5000, 0, 5000, 0, 5000]),
           ('Swedish distributor', [0, 6000, 0, 6000, 0, 6000, 0, 6000, 0, 6000, 0, 6000, 0]),
           ('Other customers', [12000, 12000, 8000, 9000, 12000, 8000, 4000, 10000, 7000, 7000, 12000, 4000, 7000]),
           ('Other receipts (VAT refunds, asset sales)', [0] * 13)]
    r = 7
    for name, vals in rec:
        ws.cell(r, 1, '   ' + name)
        for i, v in enumerate(vals):
            inp(ws.cell(r, 2 + i), v, USD)
        r += 1
    ws.cell(r, 1, 'Total receipts').font = f_bold  # r = 13
    for col in range(2, 15):
        c = ws.cell(r, col, f'=SUM({L(col)}7:{L(col)}12)'); c.number_format = USD; c.border = top_line; c.font = f_bold
    ws['A14'] = 'Payments'; ws['A14'].font = f_bold
    pay = [('Payroll', [0, 12000, 0, 12000, 0, 12000, 0, 12000, 0, 12000, 0, 12000, 0]),
           ('Materials suppliers', [15000, 15000, 15000, 15000, 15000, 15000, 15000, 15000, 15000, 15000, 15000, 15000, 15000]),
           ('Rent and utilities', [9000, 0, 0, 0, 9000, 0, 0, 0, 9000, 0, 0, 0, 9000]),
           ('Freight and export costs', [1000] * 13),
           ('Loan repayments and interest', [0, 0, 0, 6000, 0, 0, 0, 6000, 0, 0, 0, 6000, 0]),
           ('Tax payments', [0, 0, 0, 0, 0, 6000, 0, 0, 0, 5000, 0, 0, 0]),
           ('Equipment (capex)', [0] * 13),
           ('Other payments', [500] * 13)]
    r = 15
    for name, vals in pay:
        ws.cell(r, 1, '   ' + name)
        for i, v in enumerate(vals):
            inp(ws.cell(r, 2 + i), v, USD)
        r += 1
    ws.cell(23, 1, 'Total payments').font = f_bold
    for col in range(2, 15):
        c = ws.cell(23, col, f'=SUM({L(col)}15:{L(col)}22)'); c.number_format = USD; c.border = top_line; c.font = f_bold
    ws['A25'] = 'Net cash flow for the week'
    ws['A26'] = 'Closing net cash'; ws['A26'].font = f_bold
    ws['A27'] = 'Overdraft or working capital limit'
    ws['A28'] = 'Headroom (closing net cash + limit)'; ws['A28'].font = f_bold
    ws['A29'] = 'Minimum headroom you want to keep'
    ws['A30'] = 'Status'
    for col in range(2, 15):
        l = L(col)
        ws.cell(25, col, f'={l}13-{l}23').number_format = USD
        c = ws.cell(26, col, f'={l}5+{l}25'); c.number_format = USD; c.font = f_bold
        if col == 2:
            inp(ws.cell(27, col), 150000, USD); inp(ws.cell(29, col), 30000, USD)
        else:
            ws.cell(27, col, f'={L(col-1)}27').number_format = USD
            ws.cell(29, col, f'={L(col-1)}29').number_format = USD
        c = ws.cell(28, col, f'={l}26+{l}27'); c.number_format = USD; c.font = f_bold; c.border = total_line
        ws.cell(30, col, f'=IF({l}28<0,"OUT OF CASH",IF({l}28<{l}29,"Below target","OK"))').alignment = Alignment(horizontal='center')
    # totals column
    for rr in list(range(7, 14)) + list(range(15, 24)) + [25]:
        c = ws.cell(rr, 15, f'=SUM(B{rr}:N{rr})'); c.number_format = USD; c.font = f_bold
    ws['A32'] = 'Lowest headroom in the 13 weeks'; ws['A32'].font = f_bold
    c = ws['B32']; c.value = '=MIN(B28:N28)'; c.number_format = USD; c.font = f_bold
    ws['A33'] = 'Week of the lowest point'
    ws['B33'] = '=INDEX($B$4:$N$4,MATCH(MIN(B28:N28),B28:N28,0))'
    ws['A35'] = 'Example figures are BraceWorks Medical\'s (fictional), roughly matching Practice Pack B. Replace every yellow cell with your own.'; ws['A35'].font = f_sub
    ws.conditional_formatting.add('B30:N30', CellIsRule(operator='equal', formula=['"OUT OF CASH"'], fill=fill_red))
    ws.conditional_formatting.add('B30:N30', CellIsRule(operator='equal', formula=['"Below target"'], fill=fill_amber))
    ws.conditional_formatting.add('B30:N30', CellIsRule(operator='equal', formula=['"OK"'], fill=fill_green))
    ws.conditional_formatting.add('B28:N28', FormulaRule(formula=['B28<B29'], fill=fill_red))
    setw(ws, {'A': 44, **{L(c): 11 for c in range(2, 16)}})
    ws.freeze_panes = 'B5'
    # Actual vs forecast
    w2 = wb.create_sheet('Last week vs forecast')
    w2['A1'] = 'Last week: forecast vs what actually happened'; w2['A1'].font = f_title
    w2['A2'] = 'Fill this in each week before rolling the forecast forward. Big differences show which customers or costs to forecast better.'; w2['A2'].font = f_sub
    head_row(w2, 4, ['Line', 'Forecast', 'Actual', 'Difference', 'Why?'])
    lines = [('Total receipts', '=Forecast!B13', 24800), ('Total payments', '=Forecast!B23', 26100), ('Closing net cash', '=Forecast!B26', -60700)]
    for i, (n, fr, ac) in enumerate(lines):
        r = 5 + i
        w2.cell(r, 1, n)
        c = w2.cell(r, 2, fr); c.font = f_link; c.number_format = USD
        inp(w2.cell(r, 3), ac, USD)
        w2.cell(r, 4, f'=C{r}-B{r}').number_format = USD
        inp(w2.cell(r, 5), 'Example: German distributor paid $1,200 less than expected' if i == 0 else '')
    setw(w2, {'A': 22, 'B': 14, 'C': 14, 'D': 14, 'E': 60})
    base(ws); base(w2)
    instructions(wb, '13-week cash forecast', 'See your cash and borrowing headroom week by week for the next quarter, so a squeeze shows up weeks before it happens (Modules 5 and 10).',
                 ['Each week, type the opening net cash for week 1: cash in the bank minus any overdraft used.',
                  'Fill in expected receipts by customer, using the receivables aging and how each customer actually pays, not the due date.',
                  'Fill in expected payments: payroll dates, supplier due dates, rent, loan repayments, tax dates and planned purchases.',
                  'Check row 28 (headroom) and the lowest point at the bottom. Below target? Decide now what you would delay, chase or borrow.',
                  'Next week, record what actually happened on the second sheet, then delete week 1 and add a new week 13.'])
    wb.save(f'{OUT}/13-week-cash-forecast.xlsx')


# ------------------------------------------------------------------ 2. Budget builder
def budget():
    wb = Workbook(); a = wb.active; a.title = 'Assumptions'
    a['A1'] = 'Budget assumptions'; a['A1'].font = f_title
    a['A2'] = 'US dollars. Change the yellow cells; the budget P&L and cash summary update.'; a['A2'].font = f_sub
    head_row(a, 4, ['Product', 'Units last year', 'Volume growth', 'Price last year', 'Price increase', 'Material cost per unit last year', 'Material inflation', 'Budget units', 'Budget price', 'Budget material per unit', 'Budget revenue', 'Budget materials'])
    prods = [('Mattresses', 2800, 0.05, 180, 0.03, 95, 0.04), ('Slings', 7200, 0.05, 45, 0.03, 18, 0.04), ('Braces', 13000, 0.05, 28, 0.03, 11.8, 0.04)]
    for i, p in enumerate(prods):
        r = 5 + i
        a.cell(r, 1, p[0])
        inp(a.cell(r, 2), p[1], NUM); inp(a.cell(r, 3), p[2], PCT); inp(a.cell(r, 4), p[3], USD2); inp(a.cell(r, 5), p[4], PCT); inp(a.cell(r, 6), p[5], USD2); inp(a.cell(r, 7), p[6], PCT)
        a.cell(r, 8, f'=B{r}*(1+C{r})').number_format = NUM
        a.cell(r, 9, f'=D{r}*(1+E{r})').number_format = USD2
        a.cell(r, 10, f'=F{r}*(1+G{r})').number_format = USD2
        a.cell(r, 11, f'=H{r}*I{r}').number_format = USD
        a.cell(r, 12, f'=H{r}*J{r}').number_format = USD
    a['A8'] = 'Total'; a['A8'].font = f_bold
    for col in (11, 12):
        c = a.cell(8, col, f'=SUM({L(col)}5:{L(col)}7)'); c.number_format = USD; c.font = f_bold; c.border = top_line
    rows = [
        ('Other assumptions', None, None, None),
        ('Production wages last year', 152000, USD, 'Direct labour from last year\'s P&L'),
        ('Wage increase', 0.06, PCT, ''),
        ('Extra production staff', 0, NUM, ''),
        ('Cost of one extra production employee (a year)', 4000, USD, 'Including benefits'),
        ('Factory overhead last year (excluding depreciation)', 118000, USD, ''),
        ('Depreciation for the year', 45500, USD, 'Existing assets plus new equipment'),
        ('Overhead inflation', 0.04, PCT, ''),
        ('Selling and admin costs last year (excluding freight and commission)', 174000, USD, ''),
        ('Selling and admin inflation', 0.04, PCT, ''),
        ('Freight and export costs, % of sales', 0.04, PCT, ''),
        ('Sales commission, % of sales', 0.02, PCT, ''),
        ('Interest for the year', 20000, USD, ''),
        ('Tax rate', 0.25, PCT, 'Example rate; see Module 19 for your country'),
        ('Equipment purchases (capex)', 60000, USD, ''),
        ('Loan repayments', 20000, USD, ''),
        ('Receivable days (DSO)', 65, NUM, 'Target'),
        ('Inventory days', 95, NUM, 'Target'),
        ('Payable days', 63, NUM, ''),
        ('Opening receivables', 230000, USD, 'From the balance sheet'),
        ('Opening inventory', 250000, USD, ''),
        ('Opening payables', 95000, USD, ''),
        ('Opening cash', 42000, USD, ''),
    ]
    r = 10
    names = {}
    for n, v, fmt, note in rows:
        if v is None:
            a.cell(r, 1, n).font = f_bold
        else:
            a.cell(r, 1, n); inp(a.cell(r, 2), v, fmt); a.cell(r, 3, note).font = f_sub
            names[n] = f'Assumptions!$B${r}'
        r += 1
    a.cell(r + 1, 1, 'Monthly sales phasing (must total 100%)').font = f_bold
    months = ['Jan', 'Feb', 'Mar', 'Apr', 'May', 'Jun', 'Jul', 'Aug', 'Sep', 'Oct', 'Nov', 'Dec']
    phase = [78, 84, 102, 96, 98, 104, 88, 92, 112, 108, 124, 114]
    pr = r + 2
    for i, m in enumerate(months):
        a.cell(pr, 2 + i, m).font = f_bold
        inp(a.cell(pr + 1, 2 + i), round(phase[i] / 1200, 4), PCT)
    a.cell(pr + 1, 1, 'Share of the year\'s sales')
    a.cell(pr + 1, 14, f'=SUM(B{pr+1}:M{pr+1})').number_format = PCT
    a.cell(pr, 14, 'Total').font = f_bold
    a.cell(pr + 2, 14, f'=IF(ABS(N{pr+1}-1)<0.0005,"OK","Must total 100%")')
    setw(a, {'A': 58, 'B': 14, 'C': 14, 'D': 14, 'E': 14, 'F': 18, 'G': 14, 'H': 13, 'I': 13, 'J': 16, 'K': 15, 'L': 15, 'M': 10, 'N': 14})
    # P&L
    p = wb.create_sheet('Budget P&L')
    p['A1'] = 'Budget income statement by month'; p['A1'].font = f_title
    p['A2'] = 'US dollars. All formulas: change the Assumptions sheet, not this one.'; p['A2'].font = f_sub
    head_row(p, 4, ['US dollars'] + months + ['Total', '% of revenue'])
    ph = lambda col: f'Assumptions!{L(col)}${pr+1}'
    lines = [
        ('Revenue: mattresses', lambda col: f'Assumptions!$K$5*{ph(col)}'),
        ('Revenue: slings', lambda col: f'Assumptions!$K$6*{ph(col)}'),
        ('Revenue: braces', lambda col: f'Assumptions!$K$7*{ph(col)}'),
        ('Total revenue', lambda col: f'SUM({L(col)}5:{L(col)}7)'),
        ('Materials', lambda col: f'-Assumptions!$L$8*{ph(col)}'),
        ('Freight and export costs', lambda col: f'-{L(col)}8*{names["Freight and export costs, % of sales"]}'),
        ('Sales commission', lambda col: f'-{L(col)}8*{names["Sales commission, % of sales"]}'),
        ('Contribution', lambda col: f'SUM({L(col)}8:{L(col)}11)'),
        ('Production wages', lambda col: f'-({names["Production wages last year"]}*(1+{names["Wage increase"]})+{names["Extra production staff"]}*{names["Cost of one extra production employee (a year)"]}*(1+{names["Wage increase"]}))/12'),
        ('Factory overhead', lambda col: f'-{names["Factory overhead last year (excluding depreciation)"]}*(1+{names["Overhead inflation"]})/12'),
        ('Depreciation', lambda col: f'-{names["Depreciation for the year"]}/12'),
        ('Selling and admin', lambda col: f'-{names["Selling and admin costs last year (excluding freight and commission)"]}*(1+{names["Selling and admin inflation"]})/12'),
        ('Operating profit', lambda col: f'{L(col)}12+SUM({L(col)}13:{L(col)}16)'),
        ('Interest', lambda col: f'-{names["Interest for the year"]}/12'),
        ('Profit before tax', lambda col: f'{L(col)}17+{L(col)}18'),
        ('Income tax', lambda col: f'-{names["Tax rate"]}*(MAX(0,SUM($B19:{L(col)}19))-MAX(0,SUM($B19:{L(col)}19)-{L(col)}19))'),
        ('Net profit', lambda col: f'{L(col)}19+{L(col)}20'),
    ]
    bold_rows = {'Total revenue', 'Contribution', 'Operating profit', 'Profit before tax', 'Net profit'}
    for i, (n, fn) in enumerate(lines):
        r = 5 + i
        p.cell(r, 1, n).font = f_bold if n in bold_rows else f_norm
        for col in range(2, 14):
            c = p.cell(r, col, '=' + fn(col)); c.number_format = USD
            if n in bold_rows: c.font = f_bold; c.border = total_line if n == 'Net profit' else top_line
        c = p.cell(r, 14, f'=SUM(B{r}:M{r})'); c.number_format = USD; c.font = f_bold
        p.cell(r, 15, f'=IF($N$8=0,0,N{r}/$N$8)').number_format = PCT
    p['A23'] = 'Gross margin check: (revenue − materials − wages − overhead − depreciation) ÷ revenue'
    p['B23'] = '=IF(N8=0,0,(N8+N9+N13+N14+N15)/N8)'; p['B23'].number_format = PCT
    p['A24'] = 'Contribution margin'; p['B24'] = '=IF(N8=0,0,N12/N8)'; p['B24'].number_format = PCT
    p['A25'] = 'Break-even sales'; p['B25'] = '=IF(B24<=0,0,-(N13+N14+N15+N16+N18)/B24)'; p['B25'].number_format = USD
    p['A27'] = 'Income tax is worked out on year-to-date profit, so losses early in the year reduce tax later in the year. Real tax rules differ (Module 19).'; p['A27'].font = f_sub
    setw(p, {'A': 30, **{L(c): 11 for c in range(2, 14)}, 'N': 13, 'O': 12})
    p.freeze_panes = 'B5'
    # Cash
    c2 = wb.create_sheet('Cash summary')
    c2['A1'] = 'Budget cash summary for the year'; c2['A1'].font = f_title
    c2['A2'] = 'Working capital is estimated from the target days on the Assumptions sheet.'; c2['A2'].font = f_sub
    items = [
        ('Net profit', "='Budget P&L'!N21"),
        ('Add back depreciation', f'={names["Depreciation for the year"]}'),
        ('Closing receivables (DSO × revenue ÷ 365)', f"={names['Receivable days (DSO)']}*'Budget P&L'!N8/365"),
        ('Cash freed from receivables (negative = more cash tied up)', f'={names["Opening receivables"]}-B5'),
        ('Closing inventory (inventory days × materials, wages and factory overhead ÷ 365)', f"={names['Inventory days']}*-('Budget P&L'!N9+'Budget P&L'!N13+'Budget P&L'!N14)/365"),
        ('Cash freed from inventory (negative = more cash tied up)', f'={names["Opening inventory"]}-B7'),
        ('Closing payables (payable days × materials ÷ 365)', f"={names['Payable days']}*-'Budget P&L'!N9/365"),
        ('Cash from more supplier credit', f'=B9-{names["Opening payables"]}'),
        ('Cash from operations', '=B3+B4+B6+B8+B10'),
        ('Equipment purchases', f'=-{names["Equipment purchases (capex)"]}'),
        ('Loan repayments', f'=-{names["Loan repayments"]}'),
        ('Net cash flow for the year', '=B11+B12+B13'),
        ('Opening cash', f'={names["Opening cash"]}'),
        ('Closing cash (negative = extra borrowing needed)', '=B14+B15'),
    ]
    for i, (n, fml) in enumerate(items):
        r = 3 + i
        c2.cell(r, 1, n)
        c = c2.cell(r, 2, fml); c.number_format = USD
        if n.startswith(('Cash from operations', 'Net cash flow', 'Closing cash (')):
            c2.cell(r, 1).font = f_bold; c.font = f_bold; c.border = total_line if n.startswith('Closing cash (') else top_line
        if 'Closing receivables' in n or 'Closing inventory' in n or 'Closing payables' in n:
            c2.cell(r, 1).font = f_sub
    setw(c2, {'A': 62, 'B': 16})
    for w in (a, p, c2): base(w)
    instructions(wb, 'Budget builder', 'Build next year\'s budget in the right order: sales, materials, labour, overheads, then profit and cash (Module 15).',
                 ['On the Assumptions sheet, enter last year\'s units, prices and costs for each product, and your growth, price and inflation assumptions.',
                  'Enter wages, overheads, selling costs, interest, tax rate, capex and loan repayments.',
                  'Set the monthly phasing so it reflects your seasons. It must total 100%.',
                  'Read the Budget P&L for monthly profit, contribution margin and break-even, then the Cash summary for whether you need extra funding.',
                  'Save a copy as the approved budget. Use a separate copy for monthly forecasts.'],
                 ['Simplifications: production wages, overheads and selling costs are treated as fixed and spread evenly; materials, freight and commission move with sales. Tax is a simple rate on profit.'])
    wb.save(f'{OUT}/budget-builder.xlsx')


# ------------------------------------------------------------------ 3. KPI dashboard
def kpi():
    wb = Workbook(); i_ = wb.active; i_.title = 'Monthly inputs'
    months = ['Jan', 'Feb', 'Mar', 'Apr', 'May', 'Jun', 'Jul', 'Aug', 'Sep', 'Oct', 'Nov', 'Dec']
    i_['A1'] = 'Monthly inputs for the dashboard'; i_['A1'].font = f_title
    i_['A2'] = 'US dollars unless stated. Fill one column per month; the dashboard calculates the KPIs.'; i_['A2'].font = f_sub
    head_row(i_, 4, ['Input'] + months)
    data = [
        ('Revenue', [93000, 92000, 102000]),
        ('Gross profit', [25800, 25800, 27200]),
        ('EBITDA for the month', [8900, 9100, 9800]),
        ('Orders received', [96000, 98000, 106000]),
        ('Sales to largest customer', [21000, 20000, 22500]),
        ('Orders delivered on time and in full (count)', [52, 55, 58]),
        ('Orders delivered (count)', [60, 62, 65]),
        ('Material scrapped ($)', [1900, 1850, 1840]),
        ('Material issued to production ($)', [45000, 46000, 48400]),
        ('Receivables at month-end', [224000, 219000, 214000]),
        ('Inventory at month-end', [248000, 241000, 241000]),
        ('Cost of goods sold for the month', [67200, 66200, 74800]),
        ('Net debt at month-end', [170000, 162000, 154400]),
        ('EBITDA, last 12 months', [133000, 134000, 132000]),
        ('Lowest headroom in the next 13 weeks', [61000, 58000, 56000]),
    ]
    for r_, (n, vals) in enumerate(data, start=5):
        i_.cell(r_, 1, n)
        for k in range(12):
            if k < len(vals):
                inp(i_.cell(r_, 2 + k), vals[k], NUM)
            else:
                c = i_.cell(r_, 2 + k); c.fill = fill_in; c.font = f_in; c.number_format = NUM
    setw(i_, {'A': 44, **{L(c): 11 for c in range(2, 14)}})
    i_.freeze_panes = 'B5'
    d = wb.create_sheet('Dashboard')
    d['A1'] = 'Monthly KPI dashboard'; d['A1'].font = f_title
    d['A2'] = 'Status compares the latest month with the target. Off target = red, close = amber, on target = green.'; d['A2'].font = f_sub
    head_row(d, 4, ['KPI', 'Target', 'Better if', 'Owner'] + months + ['Latest', 'Status'])
    I = "'Monthly inputs'!"
    k = [
        ('Gross margin', 0.30, 'Higher', 'GM', PCT, lambda c: f'IF({I}{c}5="","",{I}{c}6/{I}{c}5)'),
        ('EBITDA margin', 0.11, 'Higher', 'GM', PCT, lambda c: f'IF({I}{c}5="","",{I}{c}7/{I}{c}5)'),
        ('Book-to-bill', 1.0, 'Higher', 'Sales', '0.00', lambda c: f'IF({I}{c}5="","",{I}{c}8/{I}{c}5)'),
        ('Largest customer share of sales', 0.20, 'Lower', 'Sales', PCT, lambda c: f'IF({I}{c}5="","",{I}{c}9/{I}{c}5)'),
        ('On time in full (OTIF)', 0.95, 'Higher', 'Production', PCT, lambda c: f'IF({I}{c}11="","",{I}{c}10/{I}{c}11)'),
        ('Scrap rate', 0.03, 'Lower', 'Production', PCT, lambda c: f'IF({I}{c}13="","",{I}{c}12/{I}{c}13)'),
        ('Days sales outstanding', 60, 'Lower', 'Finance', '0', lambda c: f'IF({I}{c}5="","",{I}{c}14/({I}{c}5*12)*365)'),
        ('Inventory days', 85, 'Lower', 'Production', '0', lambda c: f'IF({I}{c}16="","",{I}{c}15/({I}{c}16*12)*365)'),
        ('Lowest headroom, next 13 weeks', 30000, 'Higher', 'Finance', USD, lambda c: f'IF({I}{c}19="","",{I}{c}19)'),
        ('Net debt ÷ EBITDA', 2.0, 'Lower', 'GM', MULT, lambda c: f'IF({I}{c}18="","",{I}{c}17/{I}{c}18)'),
    ]
    for r_, (n, tgt, better, owner, fmt, fn) in enumerate(k, start=5):
        d.cell(r_, 1, n).font = f_bold
        inp(d.cell(r_, 2), tgt, fmt); inp(d.cell(r_, 3), better); inp(d.cell(r_, 4), owner)
        for m in range(12):
            col = L(2 + m)
            c = d.cell(r_, 5 + m, '=' + fn(col)); c.number_format = fmt
        c = d.cell(r_, 17, f'=IFERROR(LOOKUP(2,1/(E{r_}:P{r_}<>""),E{r_}:P{r_}),"")'); c.number_format = fmt; c.font = f_bold
        d.cell(r_, 18, f'=IF(Q{r_}="","",IF(C{r_}="Higher",IF(Q{r_}>=B{r_},"On target",IF(Q{r_}>=B{r_}*0.95,"Close","Off target")),IF(Q{r_}<=B{r_},"On target",IF(Q{r_}<=B{r_}*1.05,"Close","Off target"))))').alignment = Alignment(horizontal='center')
    rng = 'R5:R14'
    d.conditional_formatting.add(rng, CellIsRule(operator='equal', formula=['"Off target"'], fill=fill_red))
    d.conditional_formatting.add(rng, CellIsRule(operator='equal', formula=['"Close"'], fill=fill_amber))
    d.conditional_formatting.add(rng, CellIsRule(operator='equal', formula=['"On target"'], fill=fill_green))
    dv = DataValidation(type='list', formula1='"Higher,Lower"', allow_blank=False); d.add_data_validation(dv); dv.add('C5:C14')
    d['A16'] = 'Definitions: DSO and inventory days annualise the month (× 12). Write down any change of definition here, with the date.'; d['A16'].font = f_sub
    setw(d, {'A': 32, 'B': 10, 'C': 9, 'D': 11, **{L(c): 9 for c in range(5, 17)}, 'Q': 10, 'R': 12})
    d.freeze_panes = 'E5'
    base(i_); base(d)
    instructions(wb, 'KPI dashboard', 'A one-page monthly dashboard of 10 KPIs with targets, trend and red/amber/green status (Module 11).',
                 ['Each month, fill the next column of the Monthly inputs sheet from the monthly pack.',
                  'Set a target, direction (higher or lower is better) and owner for each KPI on the Dashboard sheet.',
                  'Read the latest column and status. Ask each owner to explain anything amber or red.',
                  'Swap KPIs to suit your business, but keep the definitions written down and unchanged.'])
    wb.move_sheet('Dashboard', offset=-1)
    wb.save(f'{OUT}/kpi-dashboard.xlsx')


# ------------------------------------------------------------------ 4. Export quote calculator
def quote():
    wb = Workbook(); q = wb.active; q.title = 'Quote'
    q['A1'] = 'Export quote calculator'; q['A1'].font = f_title
    q['A2'] = 'US dollars per unit unless stated.'; q['A2'].font = f_sub
    rows = [
        ('Product', 'Pressure-relief mattress', None),
        ('Full unit cost (materials, labour, overhead)', 125.40, USD2),
        ('Variable cost per unit (materials and other costs that rise with volume)', 95.00, USD2),
        ('Incoterm for this quote', 'CIF', None),
        ('Import duty you pay, % of price (DDP only)', 0.03, PCT),
        ('Sales commission or agent fee', 0.02, PCT),
        ('Currency buffer (if quoting in another currency)', 0.0, PCT),
        ('Customer payment terms (days)', 60, NUM),
        ('Your cost of borrowing (a year)', 0.08, PCT),
        ('Target margin after all costs', 0.25, PCT),
        ('Your list price', 180.00, USD2),
        ('Quantity on this order', 220, NUM),
    ]
    for i, (n, v, fmt) in enumerate(rows):
        r = 4 + i
        q.cell(r, 1, n); inp(q.cell(r, 2), v, fmt)
    dv = DataValidation(type='list', formula1='"EXW,FCA,FOB,CIF,DAP,DDP"'); q.add_data_validation(dv); dv.add('B7')
    q['D3'] = 'Freight, insurance and handling per unit, by Incoterm'; q['D3'].font = f_bold
    head_row(q, 4, ['Incoterm', 'Cost per unit'], start=4)
    for i, (t, v) in enumerate([('EXW', 0), ('FCA', 2), ('FOB', 3), ('CIF', 14), ('DAP', 19), ('DDP', 22)]):
        q.cell(5 + i, 4, t); inp(q.cell(5 + i, 5), v, USD2)
    q['A17'] = 'Results'; q['A17'].font = f_bold
    res = [
        ('Freight and handling for the chosen Incoterm', '=INDEX(E5:E10,MATCH(B7,D5:D10,0))', USD2),
        ('Cost of waiting to be paid, % of price', '=B12*B11/365', PCT),
        ('Duty, % of price (only if DDP)', '=IF(B7="DDP",B8,0)', PCT),
        ('All percentage costs', '=B9+B10+B19+B20', PCT),
        ('Quote price at target margin', '=IF(1-B13-B21<=0,"Check inputs",(B5+B18)/(1-B13-B21))', USD2),
        ('Price that just covers full cost (zero margin)', '=IF(1-B21<=0,"Check inputs",(B5+B18)/(1-B21))', USD2),
        ('Floor price: covers variable cost only', '=IF(1-B21<=0,"Check inputs",(B6+B18)/(1-B21))', USD2),
        ('Margin if you quote your list price', '=(B14-B5-B18-B14*B21)/B14', PCT),
        ('Order value at the quote price', '=IFERROR(B22*B15,"")', USD),
        ('Order contribution at the quote price', '=IFERROR((B22-B6-B18-B22*B21)*B15,"")', USD),
    ]
    for i, (n, fml, fmt) in enumerate(res):
        r = 18 + i
        q.cell(r, 1, n); c = q.cell(r, 2, fml); c.number_format = fmt
        if 'Quote price' in n: q.cell(r, 1).font = f_bold; c.font = f_bold; c.border = total_line
    q['A29'] = 'Never quote below the floor price. Between the floor and the zero-margin price only with spare capacity and no risk to your normal prices (Module 13).'; q['A29'].font = f_sub
    q.conditional_formatting.add('B25', CellIsRule(operator='lessThan', formula=['B13'], fill=fill_amber))
    setw(q, {'A': 66, 'B': 22, 'C': 3, 'D': 12, 'E': 14})
    base(q)
    instructions(wb, 'Export quote calculator', 'Price an export quote that covers delivery terms, payment terms, commission and currency, for a target margin (Module 12).',
                 ['Enter the product\'s full unit cost and variable cost from your costing (Module 6).',
                  'Fill the freight table with your forwarder\'s costs per unit for each Incoterm, then choose the Incoterm for this quote.',
                  'Enter commission, currency buffer, payment terms, your borrowing rate and your target margin.',
                  'Read the quote price, the zero-margin price and the floor price. Compare with your list price.'])
    wb.save(f'{OUT}/export-quote-calculator.xlsx')


# ------------------------------------------------------------------ 5. Month-end close checklist
def close():
    wb = Workbook(); c = wb.active; c.title = 'Close checklist'
    c['A1'] = 'Month-end close checklist'; c['A1'].font = f_title
    c['A2'] = 'Month:'; inp(c['B2'], 'March 2026')
    head_row(c, 4, ['Step', 'Owner', 'Due (working day)', 'Status', 'Date done', 'Evidence or notes'])
    tasks = [
        ('Bank reconciliations for every account, including foreign currency', 'Accountant', 1),
        ('Receivables ledger agrees to the balance sheet; aging produced', 'Accountant', 2),
        ('Payables ledger agrees to the balance sheet', 'Accountant', 2),
        ('Goods received not invoiced accrued', 'Accountant', 2),
        ('Freight, overtime, utilities and other accruals', 'Accountant', 3),
        ('Payroll accruals: 13th-month or bonus, contributions, leave', 'Payroll', 3),
        ('Prepayments released for the month', 'Accountant', 3),
        ('Depreciation posted', 'Accountant', 3),
        ('Stock: cycle counts reviewed, differences investigated', 'Stores and finance', 3),
        ('Warranty, returns, rebate and bad debt provisions reviewed', 'Finance manager', 4),
        ('VAT reconciled to returns; input VAT recoverable aged', 'Accountant', 4),
        ('Suspense account cleared to zero', 'Accountant', 4),
        ('Draft P&L, balance sheet and cash flow', 'Finance manager', 5),
        ('Review against budget and last month; big variances investigated', 'Finance manager', 6),
        ('Commentary written: variance, cause, action, outlook', 'Finance manager', 7),
        ('Pack issued to GM and owner', 'Finance manager', 8),
    ]
    for i, (t, o, wd) in enumerate(tasks):
        r = 5 + i
        c.cell(r, 1, t).alignment = Alignment(wrap_text=True)
        inp(c.cell(r, 2), o); inp(c.cell(r, 3), wd, '0')
        inp(c.cell(r, 4), 'Done' if i < 6 else ('In progress' if i < 9 else 'Not started'))
        inp(c.cell(r, 5), '2026-04-0' + str(min(wd, 9)) if i < 6 else ''); c.cell(r, 5).number_format = '@'
        inp(c.cell(r, 6), 'Reconciliations filed in shared folder' if i == 0 else '')
    end = 4 + len(tasks)
    dv = DataValidation(type='list', formula1='"Not started,In progress,Done"'); c.add_data_validation(dv); dv.add(f'D5:D{end}')
    c.conditional_formatting.add(f'D5:D{end}', CellIsRule(operator='equal', formula=['"Done"'], fill=fill_green))
    c.conditional_formatting.add(f'D5:D{end}', CellIsRule(operator='equal', formula=['"In progress"'], fill=fill_amber))
    c.cell(end + 2, 1, 'Steps done').font = f_bold
    c.cell(end + 2, 2, f'=COUNTIF(D5:D{end},"Done")&" of "&COUNTA(A5:A{end})')
    c.cell(end + 3, 1, 'Signed off by (name and date)').font = f_bold; inp(c.cell(end + 3, 2), '')
    setw(c, {'A': 62, 'B': 20, 'C': 16, 'D': 14, 'E': 13, 'F': 44})
    base(c)
    instructions(wb, 'Month-end close checklist', 'A fast, boring, reliable close every month: the foundation of a trustworthy monthly pack (Modules 2 and 9).',
                 ['Agree owners and due working days with your finance team, and adjust the steps to your business.',
                  'Each month, save a copy, set the month, and update the status as steps are completed.',
                  'Steps 1 to 6 come first, because everything after depends on them.',
                  'The pack shouldn\'t be issued until every step is done and signed off.'])
    wb.save(f'{OUT}/month-end-close-checklist.xlsx')


# ------------------------------------------------------------------ 6. Capex proposal
def capex():
    wb = Workbook(); s = wb.active; s.title = 'Proposal'
    s['A1'] = 'Equipment purchase proposal'; s['A1'].font = f_title
    s['A2'] = 'US dollars, before tax. Example: BraceWorks\' automatic fabric cutter (Module 14).'; s['A2'].font = f_sub
    info = [('Item', 'Automatic fabric cutting table'), ('Requested by', 'Production manager'), ('Owner of the savings', 'Production manager (fabric waste, cutting labour)'), ('Date', '2026-04-09')]
    for i, (n, v) in enumerate(info):
        s.cell(4 + i, 1, n).font = f_bold; inp(s.cell(4 + i, 2), v)
    s['A9'] = 'Investment (year 0)'; s['A9'].font = f_bold
    inv = [('Machine and software', 52000), ('Installation and electrical work', 3000), ('Validation and training', 5000), ('Extra working capital', 0)]
    for i, (n, v) in enumerate(inv):
        s.cell(10 + i, 1, '   ' + n); inp(s.cell(10 + i, 2), v, USD)
    s['A14'] = 'Total investment'; s['A14'].font = f_bold; s['B14'] = '=SUM(B10:B13)'; s['B14'].number_format = USD; s['B14'].font = f_bold; s['B14'].border = top_line
    s['A16'] = 'Yearly cash benefits and costs'; s['A16'].font = f_bold
    ben = [('Material saved', 10800), ('Labour saved (named people and dates)', 8000), ('Extra contribution from capacity', 6000), ('Running costs (maintenance, licences)', -2500)]
    for i, (n, v) in enumerate(ben):
        s.cell(17 + i, 1, '   ' + n); inp(s.cell(17 + i, 2), v, USD)
    s['A21'] = 'Net yearly benefit'; s['A21'].font = f_bold; s['B21'] = '=SUM(B17:B20)'; s['B21'].number_format = USD; s['B21'].font = f_bold; s['B21'].border = top_line
    other = [('Year 1 benefit while ramping up (% of full)', 1.0, PCT), ('Useful life (years, up to 10)', 8, '0'), ('Resale value at the end', 5000, USD), ('Discount rate (your cost of capital)', 0.10, PCT), ('Sensitivity: benefits at this % of plan', 0.70, PCT)]
    for i, (n, v, fmt) in enumerate(other):
        s.cell(23 + i, 1, n); inp(s.cell(23 + i, 2), v, fmt)
    s['A29'] = 'Cash flows'; s['A29'].font = f_bold
    head_row(s, 30, ['Year'] + [str(y) for y in range(0, 11)])
    s['A31'] = 'Plan'; s['A32'] = 'Cumulative'; s['A33'] = f'Sensitivity case'
    for y in range(0, 11):
        col = L(2 + y)
        if y == 0:
            s[f'{col}31'] = '=-B14'; s[f'{col}33'] = '=-B14'
        else:
            ramp = f'IF({y}=1,$B$23,1)'
            s[f'{col}31'] = f'=IF({y}<=$B$24,$B$21*{ramp}+IF({y}=$B$24,$B$25+$B$13,0),0)'
            s[f'{col}33'] = f'=IF({y}<=$B$24,$B$21*$B$27*{ramp}+IF({y}=$B$24,$B$25+$B$13,0),0)'
        s[f'{col}32'] = f'=SUM($B$31:{col}31)'
        for r in (31, 32, 33): s[f'{col}{r}'].number_format = USD
    s['A35'] = 'Results'; s['A35'].font = f_bold
    s['A36'] = 'Net present value (NPV)'; s['B36'] = '=B31+NPV($B$26,C31:L31)'; s['B36'].number_format = USD; s['B36'].font = f_bold
    s['A37'] = 'Internal rate of return (IRR)'; s['B37'] = '=IFERROR(IRR(B31:L31),"n/a")'; s['B37'].number_format = PCT
    s['A38'] = 'Payback (years)'; s['B38'] = '=IFERROR(MATCH(TRUE,INDEX(C32:L32>=0,0),0)-1+(-INDEX(B32:L32,MATCH(TRUE,INDEX(C32:L32>=0,0),0)))/INDEX(C31:L31,MATCH(TRUE,INDEX(C32:L32>=0,0),0)),"Never")'; s['B38'].number_format = '0.0'
    s['A39'] = 'NPV in the sensitivity case'; s['B39'] = '=B33+NPV($B$26,C33:L33)'; s['B39'].number_format = USD
    s['A40'] = 'Recommendation'; s['B40'] = '=IF(AND(B36>0,B39>0),"Robust: positive NPV even in the sensitivity case",IF(B36>0,"Positive NPV, but fragile: check the assumptions","Negative NPV: don\'t proceed as proposed"))'
    s['A42'] = 'Before approving: are labour savings named? Is validation included? Is it funded with a term loan or lease, not the overdraft? Who will do the 12-month review?'; s['A42'].font = f_sub
    s.conditional_formatting.add('B36', CellIsRule(operator='lessThan', formula=['0'], fill=fill_red))
    s.conditional_formatting.add('B39', CellIsRule(operator='lessThan', formula=['0'], fill=fill_amber))
    setw(s, {'A': 48, **{L(c): 11 for c in range(2, 13)}})
    s.column_dimensions['B'].width = 30
    base(s)
    instructions(wb, 'Equipment purchase proposal', 'A one-page case for any investment over your approval threshold, with payback, NPV, IRR and a sensitivity test (Module 14).',
                 ['Describe the item and name the person who owns the savings.',
                  'List every investment cost: price, installation, validation, training and any extra working capital.',
                  'List yearly benefits and running costs. Labour savings count only if named people or overtime really go.',
                  'Set the ramp-up, useful life, resale value, discount rate and sensitivity percentage.',
                  'Read NPV, IRR, payback and the sensitivity result. Agree a 12-month review date when approved.'])
    wb.save(f'{OUT}/equipment-purchase-proposal.xlsx')


# ------------------------------------------------------------------ 7. Annual report reader
def reader():
    wb = Workbook(); s = wb.active; s.title = 'Inputs'
    N1 = '#,##0.0;(#,##0.0);"-"'
    s['A1'] = 'Annual report reader: inputs'; s['A1'].font = f_title
    s['A2'] = 'Type figures straight from the report, in the report\'s own currency and unit. Example: Invacare Corporation, US$ millions, from its Form 10-K filings 2019 to 2022.'; s['A2'].font = f_sub
    s['A3'] = 'Company'; s['A3'].font = f_bold; inp(s['C3'], 'Invacare Corporation')
    s['A4'] = 'Currency and unit'; s['A4'].font = f_bold; inp(s['C4'], 'US$ millions')
    head_row(s, 6, ['Line', 'Where to find it', 'Year 1', 'Year 2', 'Year 3', 'Year 4'])
    years = ['2019', '2020', '2021', '2022']
    rows = [
        ('Year', '', years, None),
        ('Net sales', 'Income statement', [928.0, 850.7, 872.5, 741.7], N1),
        ('Gross profit', 'Income statement', [262.1, 245.3, 239.1, 175.4], N1),
        ('Operating profit (EBIT)', 'Income statement', [-10.4, 11.3, -24.2, -80.5], N1),
        ('Depreciation and amortisation (optional)', 'Cash flow statement, first lines', [None]*4, N1),
        ('Net profit (loss)', 'Income statement', [-53.3, -28.3, -45.6, -101.1], N1),
        ('One-off gains (+) or costs (-) inside operating profit (optional)', 'Income statement and notes', [-0.6, 9.8, -28.6, -3.3], N1),
        ('Net cash from operating activities', 'Cash flow statement', [2.7, 21.9, -14.3, -55.3], N1),
        ('Capital spending (type as a positive number)', 'Cash flow statement, investing', [10.9, 22.3, 17.7, 3.8], N1),
        ('Cash and cash equivalents', 'Balance sheet', [80.1, 105.3, 83.7, 58.8], N1),
        ('Trade receivables', 'Balance sheet', [116.7, 108.6, 117.1, 88.0], N1),
        ('Inventories', 'Balance sheet', [120.5, 115.5, 144.3, 112.6], N1),
        ('Trade payables (optional)', 'Balance sheet', [None]*4, N1),
        ('Total borrowings', 'Balance sheet or debt note', [219.5, 245.1, 308.1, 354.2], N1),
        ('Total equity', 'Balance sheet', [308.5, 333.8, 218.5, 81.1], N1),
        ('Total assets (optional)', 'Balance sheet', [None]*4, N1),
        ('Goodwill', 'Balance sheet or goodwill note', [373.4, 402.5, 355.9, 326.3], N1),
    ]
    for i, (lab, where, vals, fmt) in enumerate(rows):
        r = 7 + i
        s.cell(r, 1, lab).font = f_bold if i == 0 else f_norm
        s.cell(r, 2, where).font = f_sub
        for j, v in enumerate(vals):
            c = s.cell(r, 3 + j)
            inp(c, v, fmt)
    s['A25'] = 'Invacare notes: one-off items are impairments (2019, 2021, 2022) and the gain on selling Dynamic Controls (2020). Restructuring costs came every year, so they are left in. Borrowings exclude finance leases, as in the debt note. Source: Invacare Form 10-K filings on SEC EDGAR.'
    s['A25'].font = f_sub; s['A25'].alignment = Alignment(wrap_text=True); s.merge_cells('A25:F25'); s.row_dimensions[25].height = 42
    s['A26'] = 'Leave optional lines blank if the report doesn\'t give them: the ratios that need them will show n/a.'; s['A26'].font = f_sub
    s.cell(8, 3).comment = Comment('Net sales, first line of the income statement. Example from Invacare 10-K, FY2019.', 'GM Ledger')
    setw(s, {'A': 52, 'B': 30, 'C': 13, 'D': 13, 'E': 13, 'F': 13})
    s.freeze_panes = 'C7'
    base(s)

    # Ratios and flags
    r_ = wb.create_sheet('Ratios and flags')
    r_['A1'] = 'Ratios and warning signs'; r_['A1'].font = f_title
    r_['A2'] = '="Company: "&Inputs!C3&"   ·   "&Inputs!C4'; r_['A2'].font = f_sub
    head_row(r_, 4, ['Ratio', 'What it tells you', 'Year 1', 'Year 2', 'Year 3', 'Year 4'])
    I = lambda row, col: f"Inputs!{col}{row}"
    cols = ['C', 'D', 'E', 'F']
    def prev(c): return cols[cols.index(c) - 1]
    lines = [
        ('Year', '', lambda c: f"=Inputs!{c}7", None, False),
        ('Sales growth', 'Is the business growing?', lambda c: f'=IF(OR(Inputs!{prev(c)}8="",Inputs!{prev(c)}8=0),"n/a",Inputs!{c}8/Inputs!{prev(c)}8-1)', PCT, True),
        ('Gross margin', 'Pricing and factory efficiency', lambda c: f'=IF(Inputs!{c}8=0,"n/a",Inputs!{c}9/Inputs!{c}8)', PCT, False),
        ('Operating margin', 'Profit from operations per sales dollar', lambda c: f'=IF(Inputs!{c}8=0,"n/a",Inputs!{c}10/Inputs!{c}8)', PCT, False),
        ('Underlying operating margin (without one-offs)', 'What the ongoing business earned', lambda c: f'=IF(Inputs!{c}8=0,"n/a",(Inputs!{c}10-N(Inputs!{c}13))/Inputs!{c}8)', PCT, False),
        ('EBITDA', 'Needs depreciation and amortisation', lambda c: f'=IF(Inputs!{c}11="","n/a",Inputs!{c}10+Inputs!{c}11)', '#,##0.0;(#,##0.0);"-"', False),
        ('Net margin', 'What owners keep', lambda c: f'=IF(Inputs!{c}8=0,"n/a",Inputs!{c}12/Inputs!{c}8)', PCT, False),
        ('Cash conversion (operating cash flow ÷ EBITDA)', 'Is profit turning into cash?', lambda c: f'=IF(OR(Inputs!{c}11="",Inputs!{c}10+Inputs!{c}11<=0),"n/a",Inputs!{c}14/(Inputs!{c}10+Inputs!{c}11))', PCT, False),
        ('Free cash flow (operating cash flow − capital spending)', 'Can the business pay its own way?', lambda c: f'=Inputs!{c}14-Inputs!{c}15', '#,##0.0;(#,##0.0);"-"', False),
        ('Free cash flow % of sales', '', lambda c: f'=IF(Inputs!{c}8=0,"n/a",(Inputs!{c}14-Inputs!{c}15)/Inputs!{c}8)', PCT, False),
        ('Days sales outstanding', 'How fast customers pay', lambda c: f'=IF(Inputs!{c}8=0,"n/a",Inputs!{c}17/Inputs!{c}8*365)', '0', False),
        ('Inventory days', 'How long stock sits', lambda c: f'=IF(Inputs!{c}8-Inputs!{c}9<=0,"n/a",Inputs!{c}18/(Inputs!{c}8-Inputs!{c}9)*365)', '0', False),
        ('Payable days', 'How long it takes to pay suppliers', lambda c: f'=IF(OR(Inputs!{c}19="",Inputs!{c}8-Inputs!{c}9<=0),"n/a",Inputs!{c}19/(Inputs!{c}8-Inputs!{c}9)*365)', '0', False),
        ('Net debt (borrowings − cash)', '', lambda c: f'=Inputs!{c}20-Inputs!{c}16', '#,##0.0;(#,##0.0);"-"', False),
        ('Net debt ÷ EBITDA', 'Years of EBITDA to repay debt', lambda c: f'=IF(OR(Inputs!{c}11="",Inputs!{c}10+Inputs!{c}11<=0),"n/a",(Inputs!{c}20-Inputs!{c}16)/(Inputs!{c}10+Inputs!{c}11))', MULT, False),
        ('Debt ÷ equity', 'Lenders\' money vs owners\' money', lambda c: f'=IF(Inputs!{c}21<=0,"n/a",Inputs!{c}20/Inputs!{c}21)', MULT, False),
        ('Goodwill ÷ equity', 'Could a write-down wipe out equity?', lambda c: f'=IF(Inputs!{c}21<=0,"n/a",Inputs!{c}23/Inputs!{c}21)', MULT, False),
        ('Goodwill ÷ total assets', 'Needs total assets', lambda c: f'=IF(OR(Inputs!{c}22="",Inputs!{c}22=0),"n/a",Inputs!{c}23/Inputs!{c}22)', PCT, False),
    ]
    for i, (lab, why, fn, fmt, skip1) in enumerate(lines):
        r = 5 + i
        r_.cell(r, 1, lab).font = f_bold if i == 0 else f_norm
        r_.cell(r, 2, why).font = f_sub
        for c in cols:
            if skip1 and c == 'C':
                r_[f'{c}{r}'] = 'n/a'; continue
            cell = r_[f'{c}{r}']; cell.value = fn(c); cell.font = f_link if i == 0 else f_norm
            if fmt: cell.number_format = fmt
            cell.alignment = Alignment(horizontal='right')
    fr = 5 + len(lines) + 1
    r_.cell(fr, 1, 'Warning signs (Module 21)').font = f_bold
    head_row(r_, fr + 1, ['Warning sign', 'Why it matters', 'Year 1', 'Year 2', 'Year 3', 'Year 4'])
    flags = [
        ('Free cash flow negative', 'The business is spending more than it generates', lambda c: f'=IF(Inputs!{c}14-Inputs!{c}15<0,"Yes","-")', False),
        ('Operating cash flow negative', 'Day-to-day operations use cash', lambda c: f'=IF(Inputs!{c}14<0,"Yes","-")', False),
        ('Receivables grew faster than sales', 'Customers paying slower, or sales quality falling', lambda c: f'=IF(OR(Inputs!{prev(c)}17="",Inputs!{prev(c)}17=0,Inputs!{prev(c)}8=0),"-",IF(Inputs!{c}17/Inputs!{prev(c)}17>Inputs!{c}8/Inputs!{prev(c)}8+0.02,"Yes","-"))', True),
        ('Inventory grew faster than sales', 'Cash tied up in stock; possible slow-moving stock', lambda c: f'=IF(OR(Inputs!{prev(c)}18="",Inputs!{prev(c)}18=0,Inputs!{prev(c)}8=0),"-",IF(Inputs!{c}18/Inputs!{prev(c)}18>Inputs!{c}8/Inputs!{prev(c)}8+0.02,"Yes","-"))', True),
        ('Gross margin fell more than 1 point', 'Pricing or cost pressure', lambda c: f'=IF(OR(Inputs!{c}8=0,Inputs!{prev(c)}8=0),"-",IF(Inputs!{c}9/Inputs!{c}8<Inputs!{prev(c)}9/Inputs!{prev(c)}8-0.01,"Yes","-"))', True),
        ('Net debt rose', 'More reliance on lenders', lambda c: f'=IF((Inputs!{c}20-Inputs!{c}16)>(Inputs!{prev(c)}20-Inputs!{prev(c)}16)+0.005*ABS(Inputs!{prev(c)}20),"Yes","-")', True),
        ('Equity fell by more than the loss', 'Other losses going straight to equity: read the statement of changes in equity', lambda c: f'=IF((Inputs!{prev(c)}21-Inputs!{c}21)-MAX(0,-Inputs!{c}12)>0.02*ABS(Inputs!{prev(c)}21),"Yes","-")', True),
        ('Goodwill larger than equity', 'A write-down could wipe out the owners\' cushion', lambda c: f'=IF(Inputs!{c}23>Inputs!{c}21,"Yes","-")', False),
        ('One-offs over half of operating profit', 'Headline profit not representative', lambda c: f'=IF(Inputs!{c}13="","-",IF(ABS(Inputs!{c}13)>0.5*ABS(Inputs!{c}10),"Yes","-"))', False),
    ]
    for i, (lab, why, fn, skip1) in enumerate(flags):
        r = fr + 2 + i
        r_.cell(r, 1, lab); r_.cell(r, 2, why).font = f_sub
        for c in cols:
            if skip1 and c == 'C':
                r_[f'{c}{r}'] = '-'; r_[f'{c}{r}'].alignment = Alignment(horizontal='center'); continue
            r_[f'{c}{r}'] = fn(c); r_[f'{c}{r}'].alignment = Alignment(horizontal='center')
    last = fr + 1 + len(flags)
    tr = last + 1
    r_.cell(tr, 1, 'Warning signs this year').font = f_bold
    for c in cols:
        r_[f'{c}{tr}'] = f'=COUNTIF({c}{fr+2}:{c}{last},"Yes")'; r_[f'{c}{tr}'].font = f_bold; r_[f'{c}{tr}'].border = top_line; r_[f'{c}{tr}'].alignment = Alignment(horizontal='center')
    r_.conditional_formatting.add(f'C{fr+2}:F{last}', CellIsRule(operator='equal', formula=['"Yes"'], fill=fill_red))
    r_.conditional_formatting.add(f'C{tr}:F{tr}', CellIsRule(operator='greaterThanOrEqual', formula=['4'], fill=fill_red))
    r_.conditional_formatting.add(f'C{tr}:F{tr}', CellIsRule(operator='between', formula=['2', '3'], fill=fill_amber))
    r_.cell(tr + 2, 1, 'Signs are prompts to read further, not verdicts. Check each one against the notes and the cash flow statement.').font = f_sub
    setw(r_, {'A': 50, 'B': 46, 'C': 12, 'D': 12, 'E': 12, 'F': 12})
    r_.freeze_panes = 'C5'
    base(r_)
    ws = instructions(wb, 'Annual report reader', 'Turn any company\'s annual report into ratios and warning signs in ten minutes: a competitor, a customer, a supplier, or your own audited accounts (Module 21).',
                 ['Type the company name and the report\'s currency and unit on the Inputs sheet.',
                  'Copy up to four years of figures from the statements into the yellow cells. Start with the oldest year in column C.',
                  'Leave optional lines blank if the report doesn\'t give them.',
                  'Open Ratios and flags. Read each red warning sign, then find the note or statement that explains it.',
                  'Write three sentences: how the company is doing, its biggest risk, and what you\'d watch next year.'])
    ws['A2'] = 'The GM\'s Ledger course template. The example figures are Invacare Corporation\'s, from its published Form 10-K filings (US$ millions): replace them with any company\'s.'
    wb.save(f'{OUT}/annual-report-reader.xlsx')


# ------------------------------------------------------------------ 8. Three-statement forecast
def forecast3():
    wb = Workbook(); a = wb.active; a.title = 'Assumptions'
    a['A1'] = 'Three-statement forecast: assumptions'; a['A1'].font = f_title
    a['A2'] = 'US dollars. Change only the yellow cells. Example: BraceWorks Medical (fictional), opening from its 31 December balance sheet (Module 1).'; a['A2'].font = f_sub
    head_row(a, 4, ['Driver', 'Year 1', 'Year 2', 'Year 3', 'Note'])
    drv = [('Sales growth', [0.05, 0.05, 0.05], PCT, 'Module 15'),
           ('Gross margin before depreciation', [0.315, 0.315, 0.315], PCT, 'Last year: $378,000 ÷ $1,200,000'),
           ('Selling and admin cost increase', [0.04, 0.04, 0.04], PCT, 'Wages and inflation'),
           ('Customer days (DSO)', [70, 70, 70], '0', 'Module 16'),
           ('Stock days (of cost of sales incl. depreciation)', [106, 106, 106], '0', 'Module 7'),
           ('Supplier days (of materials)', [63, 63, 63], '0', 'Module 16'),
           ('Materials as % of revenue', [0.46, 0.46, 0.46], PCT, 'For supplier days'),
           ('Equipment purchases', [60000, 30000, 30000], USD, 'Module 14'),
           ('New equipment loan', [0, 0, 0], USD, 'Module 20'),
           ('Loan repayments', [20000, 20000, 20000], USD, 'Loan schedule'),
           ('Dividends', [0, 20000, 20000], USD, 'Module 20')]
    for i, (n, vals, fmt, note) in enumerate(drv):
        r = 5 + i
        a.cell(r, 1, n)
        for j, v in enumerate(vals):
            inp(a.cell(r, 2 + j), v, fmt)
        a.cell(r, 5, note).font = f_sub
    a['A17'] = 'Fixed rules'; a['A17'].font = f_bold
    fixed = [('Interest rate on overdraft and loans', 0.064, PCT, 'Charged on opening balances, plus half a year on new loans'),
             ('Tax rate', 0.25, PCT, 'Example rate: see Module 19 for your country'),
             ('Useful life of new equipment (years)', 8, '0', 'Half a year of depreciation in the year bought'),
             ('Depreciation on existing equipment, per year', 38000, USD, 'Last year\'s charge'),
             ('Minimum cash balance', 20000, USD, 'The overdraft covers any shortfall below this'),
             ('Overdraft limit', 150000, USD, 'From the bank facility letter'),
             ('Covenant: net debt ÷ EBITDA no more than', 2.5, MULT, 'From the loan agreement')]
    for i, (n, v, fmt, note) in enumerate(fixed):
        r = 18 + i
        a.cell(r, 1, n); inp(a.cell(r, 2), v, fmt); a.cell(r, 5, note).font = f_sub
    a['A26'] = 'Last year and opening balance sheet (31 December)'; a['A26'].font = f_bold
    openb = [('Revenue last year', 1200000), ('Selling and admin last year', 246000),
             ('Cash', 42000), ('Receivables', 230000), ('Inventory', 250000), ('Prepayments', 13000), ('Equipment, net', 310000),
             ('Payables', 95000), ('Accrued expenses', 30000), ('Tax payable', 8500), ('Overdraft', 120000), ('Equipment loan', 100000),
             ('Share capital', 250000), ('Retained earnings', 241500)]
    for i, (n, v) in enumerate(openb):
        r = 27 + i
        a.cell(r, 1, n); inp(a.cell(r, 2), v, USD)
    a['A41'] = 'Opening balance check (must be zero)'; a['A41'].font = f_bold
    a['B41'] = '=SUM(B29:B33)-SUM(B34:B40)'; a['B41'].number_format = USD
    a.conditional_formatting.add('B41', CellIsRule(operator='notEqual', formula=['0'], fill=fill_red))
    setw(a, {'A': 48, 'B': 14, 'C': 14, 'D': 14, 'E': 52})
    base(a)

    f = wb.create_sheet('Forecast')
    f['A1'] = 'Three-statement forecast'; f['A1'].font = f_title
    f['A2'] = 'US dollars. All formulas: change the Assumptions sheet, not this one. Negative numbers in brackets.'; f['A2'].font = f_sub
    head_row(f, 4, ['', 'Last year / opening', 'Year 1', 'Year 2', 'Year 3'])
    Y = ['C', 'D', 'E']; A = {'C': 'B', 'D': 'C', 'E': 'D'}  # forecast column -> assumptions column
    P = {'C': 'B', 'D': 'C', 'E': 'D'}  # previous column on Forecast
    def link(c):
        c.font = f_link
    rows = {}
    r = 5
    def put(label, fn, fmt=USD, bold=False, opening=None, border=None):
        nonlocal r
        f.cell(r, 1, label).font = f_bold if bold else f_norm
        if opening is not None:
            c = f.cell(r, 2, opening); c.number_format = fmt
            if isinstance(opening, str) and 'Assumptions' in opening: link(c)
        for col in Y:
            c = f[f'{col}{r}']; c.value = '=' + fn(col); c.number_format = fmt
            if bold: c.font = f_bold
            if border: c.border = border
        rows[label] = r; r += 1
        return r - 1
    def sec(t):
        nonlocal r
        f.cell(r, 1, t).font = f_bold; f.cell(r, 1).fill = fill_band
        for col in ['B'] + Y: f[f'{col}{r}'].fill = fill_band
        r += 1
    sec('Income statement')
    rv = put('Revenue', lambda c: f"{P[c]}{r}*(1+Assumptions!{A[c]}5)", opening='=Assumptions!B27')
    cs = put('Cost of sales (excl. depreciation)', lambda c: f"-{c}{rv}*(1-Assumptions!{A[c]}6)")
    sg = put('Selling and admin', lambda c: f"{P[c]}{r}*(1+Assumptions!{A[c]}7)" if c != 'C' else f"-Assumptions!B28*(1+Assumptions!B7)", opening=None)
    eb = put('EBITDA', lambda c: f"{c}{rv}+{c}{cs}+{c}{sg}", bold=True, border=top_line)
    dp = put('Depreciation', lambda c: {'C': "-(Assumptions!$B$21+Assumptions!B12/(2*Assumptions!$B$20))",
                                        'D': "-(Assumptions!$B$21+Assumptions!B12/Assumptions!$B$20+Assumptions!C12/(2*Assumptions!$B$20))",
                                        'E': "-(Assumptions!$B$21+(Assumptions!B12+Assumptions!C12)/Assumptions!$B$20+Assumptions!D12/(2*Assumptions!$B$20))"}[c])
    op = put('Operating profit', lambda c: f"{c}{eb}+{c}{dp}", bold=True, border=top_line)
    it_row = r; r += 1   # interest filled after balance sheet rows exist
    pb = put('Profit before tax', lambda c: f"{c}{op}+{c}{it_row}", bold=True, border=top_line)
    tx = put('Tax', lambda c: f"-MAX(0,{c}{pb})*Assumptions!$B$19")
    npr = put('Net profit', lambda c: f"{c}{pb}+{c}{tx}", bold=True, border=total_line)
    r += 1
    sec('Balance sheet, 31 December')
    cash = r; r += 1
    ar = put('Receivables', lambda c: f"{c}{rv}*Assumptions!{A[c]}8/365", opening='=Assumptions!B30')
    iv = put('Inventory', lambda c: f"-({c}{cs}+{c}{dp})*Assumptions!{A[c]}9/365", opening='=Assumptions!B31')
    pp = put('Prepayments', lambda c: f"{P[c]}{r}", opening='=Assumptions!B32')
    eq_ = put('Equipment, net', lambda c: f"{P[c]}{r}+Assumptions!{A[c]}12+{c}{dp}", opening='=Assumptions!B33')
    ta = put('Total assets', lambda c: f"SUM({c}{cash}:{c}{eq_})", bold=True, border=top_line, opening=f'=SUM(B{cash}:B{eq_})')
    ap = put('Payables', lambda c: f"{c}{rv}*Assumptions!{A[c]}11*Assumptions!{A[c]}10/365", opening='=Assumptions!B34')
    ac = put('Accrued expenses', lambda c: f"{P[c]}{r}", opening='=Assumptions!B35')
    tp = put('Tax payable', lambda c: f"{P[c]}{r}", opening='=Assumptions!B36')
    od = r; r += 1
    ln = put('Equipment loan', lambda c: f"{P[c]}{r}-MIN(Assumptions!{A[c]}14,{P[c]}{r})+Assumptions!{A[c]}13", opening='=Assumptions!B38')
    sc = put('Share capital', lambda c: f"{P[c]}{r}", opening='=Assumptions!B39')
    re_ = put('Retained earnings', lambda c: f"{P[c]}{r}+{c}{npr}-Assumptions!{A[c]}15", opening='=Assumptions!B40')
    tl = put('Total liabilities and equity', lambda c: f"SUM({c}{ap}:{c}{re_})", bold=True, border=total_line, opening=f'=SUM(B{ap}:B{re_})')
    r += 1
    sec('Cash flow statement')
    c1 = put('Net profit', lambda c: f"{c}{npr}")
    c2 = put('Add back depreciation', lambda c: f"-{c}{dp}")
    c3 = put('Receivables (increase) / decrease', lambda c: f"{P[c]}{ar}-{c}{ar}")
    c4 = put('Inventory (increase) / decrease', lambda c: f"{P[c]}{iv}-{c}{iv}")
    c5 = put('Payables increase / (decrease)', lambda c: f"{c}{ap}-{P[c]}{ap}")
    cfo = put('Cash from operations', lambda c: f"SUM({c}{c1}:{c}{c5})", bold=True, border=top_line)
    cx = put('Equipment purchases', lambda c: f"-Assumptions!{A[c]}12")
    fcf = put('Free cash flow', lambda c: f"{c}{cfo}+{c}{cx}", bold=True, border=top_line)
    rp = put('Loan repayments', lambda c: f"-MIN(Assumptions!{A[c]}14,{P[c]}{ln})")
    nl = put('New equipment loan', lambda c: f"Assumptions!{A[c]}13")
    dv = put('Dividends', lambda c: f"-Assumptions!{A[c]}15")
    pre = put('Cash before overdraft', lambda c: f"{P[c]}{cash}+{c}{fcf}+{c}{rp}+{c}{nl}+{c}{dv}", bold=True, border=top_line)
    odc = put('Overdraft drawn / (repaid)', lambda c: f"IF({c}{pre}<Assumptions!$B$22,Assumptions!$B$22-{c}{pre},-MIN({P[c]}{od},{c}{pre}-Assumptions!$B$22))")
    cc = put('Closing cash', lambda c: f"{c}{pre}+{c}{odc}", bold=True, border=total_line)
    # deferred rows: interest, cash, overdraft
    f.cell(it_row, 1, 'Interest')
    for c in Y:
        f[f'{c}{it_row}'] = f"=-Assumptions!$B$18*({P[c]}{od}+{P[c]}{ln}+Assumptions!{A[c]}13/2)"; f[f'{c}{it_row}'].number_format = USD
    f.cell(cash, 1, 'Cash'); f.cell(cash, 2, '=Assumptions!B29').number_format = USD; link(f.cell(cash, 2))
    for c in Y:
        f[f'{c}{cash}'] = f'={c}{cc}'; f[f'{c}{cash}'].number_format = USD
    f.cell(od, 1, 'Overdraft'); f.cell(od, 2, '=Assumptions!B37').number_format = USD; link(f.cell(od, 2))
    for c in Y:
        f[f'{c}{od}'] = f'={P[c]}{od}+{c}{odc}'; f[f'{c}{od}'].number_format = USD
    for rr in (rv, ar, iv, pp, eq_, ac, tp, ln, sc, re_):
        if f.cell(rr, 2).value: link(f.cell(rr, 2))
    f.cell(rv, 2).value = '=Assumptions!B27'; f.cell(sg, 2).value = '=-Assumptions!B28'; f.cell(sg, 2).number_format = USD; link(f.cell(sg, 2))
    # fix selling & admin year 2/3 sign (previous is negative)
    for c in ['D', 'E']:
        f[f'{c}{sg}'] = f"={P[c]}{sg}*(1+Assumptions!{A[c]}7)"
    r += 1
    sec('Checks and limits')
    chk = put('Balance check: assets − (liabilities + equity)', lambda c: f"ROUND({c}{ta}-{c}{tl},0)", bold=True)
    f.cell(chk, 2, f'=ROUND(B{ta}-B{tl},0)').number_format = USD
    hd = put('Overdraft headroom', lambda c: f"Assumptions!$B$23-{c}{od}")
    nd = put('Net debt ÷ EBITDA', lambda c: f"IF({c}{eb}<=0,\"n/a\",({c}{od}+{c}{ln}-{c}{cash})/{c}{eb})", fmt=MULT)
    st = put('Status', lambda c: f"IF({c}{chk}<>0,\"Check the model\",IF({c}{hd}<0,\"Over the overdraft limit\",IF(AND(ISNUMBER({c}{nd}),{c}{nd}>Assumptions!$B$24),\"Covenant breached\",IF({c}{eb}<=0,\"Covenant breached\",\"OK\"))))", fmt='General', bold=True)
    f.conditional_formatting.add(f'C{chk}:E{chk}', CellIsRule(operator='notEqual', formula=['0'], fill=fill_red))
    f.conditional_formatting.add(f'C{hd}:E{hd}', CellIsRule(operator='lessThan', formula=['0'], fill=fill_red))
    f.conditional_formatting.add(f'C{st}:E{st}', CellIsRule(operator='equal', formula=['"OK"'], fill=fill_green))
    f.conditional_formatting.add(f'C{st}:E{st}', CellIsRule(operator='notEqual', formula=['"OK"'], fill=fill_red))
    setw(f, {'A': 46, 'B': 18, 'C': 14, 'D': 14, 'E': 14})
    f.freeze_panes = 'B5'
    base(f)
    instructions(wb, 'Three-statement forecast', 'A linked three-year forecast of the income statement, balance sheet and cash flow, with the overdraft as the balancing item, covenant tests and a balance check (Module 22).',
                 ['On the Assumptions sheet, type your opening balance sheet from last year\'s accounts. The opening balance check must read zero.',
                  'Set the drivers for each year: growth, margin, cost increases, customer, stock and supplier days, equipment, loans and dividends.',
                  'Check the fixed rules: interest rate, tax rate, minimum cash, overdraft limit and covenant.',
                  'Open the Forecast sheet. The balance check must read zero; the status row shows whether a limit is breached.',
                  'Save a copy, then change a few drivers for a downside case and compare. Know what breaks first.'])
    wb.save(f'{OUT}/three-statement-forecast.xlsx')

for fn in (cash13, budget, kpi, quote, close, capex, reader, forecast3):
    fn()
print('written', sorted(os.listdir(OUT)))
